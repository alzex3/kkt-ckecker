import asyncio
import logging
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from client import AsyncKKTChecker, CheckResultEnum
from config import Config


@dataclass
class ProcessingProgress:
    current: int
    total: int
    current_item: str

    @property
    def percentage(self) -> float:
        return (self.current / self.total * 100) if self.total > 0 else 0

    def __str__(self) -> str:
        return f"{self.current}/{self.total} - {self.percentage:.1f}%"


class ExcelProcessor:
    def __init__(self, config: Config) -> None:
        self.config = config
        self.checker = AsyncKKTChecker()

    async def process_file(
        self,
        input_path: Path,
        output_path: Path,
        progress_callback: Callable[[ProcessingProgress], None] | None = None,
    ) -> None:
        try:
            wb = load_workbook(str(input_path))

            if "ККТ" not in wb.sheetnames:
                raise ValueError("Worksheet 'ККТ' not found in the Excel file")

            ws = wb["ККТ"]

            # Find column indices
            title_column_index = None
            model_column_index = None
            number_column_index = None
            state_column_index = None

            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            for i, cell_value in enumerate(header):
                match cell_value:
                    case "Наименование":
                        title_column_index = i
                    case "Модель ККТ":
                        model_column_index = i
                    case "Заводской Номер ККТ":
                        number_column_index = i
                    case "Статус Проверки":
                        state_column_index = i

            if None in (title_column_index, model_column_index, number_column_index, state_column_index):
                raise ValueError("Required columns not found in the Excel file")

            # Count valid rows
            rows_to_process = []
            for row in ws.iter_rows(min_row=2, max_row=400):
                owner_cell = row[title_column_index]
                model_cell = row[model_column_index]
                number_cell = row[number_column_index]

                if owner_cell.value and model_cell.value and number_cell.value:
                    rows_to_process.append((row, owner_cell, model_cell, number_cell))

            total_rows = len(rows_to_process)
            if total_rows == 0:
                raise ValueError("No valid rows found to process")

            # Process rows with concurrency control
            semaphore = asyncio.Semaphore(self.config.max_concurrent_checks)
            processed_count = 0

            async def process_row(row_data):
                nonlocal processed_count
                row, owner_cell, model_cell, number_cell = row_data
                state_cell = row[state_column_index]

                async with semaphore:
                    try:

                        def update_progress():
                            if progress_callback:
                                progress = ProcessingProgress(
                                    current=processed_count + 1,
                                    total=total_rows,
                                    current_item=f"{owner_cell.value}",
                                )
                                # Handle both sync and async callbacks
                                if asyncio.iscoroutinefunction(progress_callback):
                                    asyncio.create_task(progress_callback(progress))
                                else:
                                    progress_callback(progress)

                        update_progress()

                        result = await self.checker.check(
                            model=str(model_cell.value),
                            factory_number=str(number_cell.value),
                            progress_callback=lambda msg: update_progress(),
                        )

                        # Update cell based on result
                        match result:
                            case CheckResultEnum.FAILED:
                                state_cell.value = "НЕУСПЕШНО"
                                state_cell.fill = PatternFill(
                                    start_color="FF0000", end_color="FF0000", fill_type="solid"
                                )
                            case CheckResultEnum.SUCCESS:
                                state_cell.value = "УСПЕШНО"
                                state_cell.fill = PatternFill(
                                    start_color="00FF00", end_color="00FF00", fill_type="solid"
                                )
                            case CheckResultEnum.ERROR:
                                state_cell.value = "ОШИБКА"
                                state_cell.fill = PatternFill(
                                    start_color="FFFF00", end_color="FFFF00", fill_type="solid"
                                )

                        processed_count += 1
                        update_progress()

                        # Add delay to avoid overwhelming the API
                        await asyncio.sleep(self.config.api_delay_seconds)

                    except Exception as e:
                        logging.error(f"Error processing row {processed_count + 1}: {e}")
                        state_cell.value = "ОШИБКА"
                        state_cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
                        processed_count += 1

            # Process all rows concurrently
            tasks = [process_row(row_data) for row_data in rows_to_process]
            await asyncio.gather(*tasks)

            # Save the workbook
            wb.save(str(output_path))

            if progress_callback:
                final_progress = ProcessingProgress(
                    current=total_rows,
                    total=total_rows,
                    current_item="",
                )
                # Handle both sync and async callbacks
                if asyncio.iscoroutinefunction(progress_callback):
                    await progress_callback(final_progress)
                else:
                    progress_callback(final_progress)

        except Exception as e:
            logging.error(f"Error processing Excel file: {e}")
            raise
        finally:
            await self.checker.close()

    @staticmethod
    async def validate_file(file_path: Path) -> bool:
        try:
            wb = load_workbook(str(file_path))

            if "ККТ" not in wb.sheetnames:
                return False

            ws = wb["ККТ"]
            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))

            required_columns = {"Наименование", "Модель ККТ", "Заводской Номер ККТ", "Статус Проверки"}
            found_columns = {cell for cell in header if cell in required_columns}

            return len(found_columns) == len(required_columns)

        except Exception as e:
            logging.error(f"Error validating file: {e}")
            return False
